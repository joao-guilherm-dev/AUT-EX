import os
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
import aut  # noqa: E402


@pytest.mark.parametrize(
    "texto, tipo, esperado",
    [("10", int, 10), ("10.5", float, 10.5), ("10,5", float, 10.5), ("1.500,50", float, 1500.5), (" 3 ", int, 3)],
)
def test_converter_numero_valido(texto, tipo, esperado):
    assert aut.converter_numero(texto, tipo, 0) == esperado


@pytest.mark.parametrize("texto", ["abc", "", "nan", "inf", "-1"])
def test_converter_numero_invalido(texto):
    with pytest.raises(ValueError):
        aut.converter_numero(texto, float, 0)


def test_quantidade_nao_aceita_decimal_nem_zero():
    with pytest.raises(ValueError):
        aut.converter_numero("2.5", int, 1)
    with pytest.raises(ValueError):
        aut.converter_numero("0", int, 1)


def test_planilha_nova_tem_cabecalho(tmp_path):
    _, ws = aut.abrir_planilha(tmp_path / "x.xlsx")
    assert [c.value for c in ws[1]] == aut.CABECALHO
    assert ws.title == aut.NOME_ABA


def test_total_e_formula_do_excel(tmp_path):
    _, ws = aut.abrir_planilha(tmp_path / "x.xlsx")
    linha = aut.adicionar_venda(ws, "Mouse", 2, 49.9)
    assert linha == 2
    assert ws["D2"].value == "=B2*C2"


def test_formatacao_nao_atrapalha_proxima_linha(tmp_path):
    _, ws = aut.abrir_planilha(tmp_path / "x.xlsx")
    aut.adicionar_venda(ws, "Mouse", 1, 10)
    aut.formatar_planilha(ws)  # escreve o resumo em F1:G2
    assert aut.adicionar_venda(ws, "Teclado", 1, 20) == 3


def test_formatacao(tmp_path):
    _, ws = aut.abrir_planilha(tmp_path / "x.xlsx")
    aut.adicionar_venda(ws, "Monitor", 1, 900)
    aut.formatar_planilha(ws)
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:D2"
    assert ws["C2"].number_format == aut.FORMATO_MOEDA
    assert ws["A1"].font.bold
    assert ws["G1"].value == "=SUM(D:D)"


def test_fluxo_completo_e_reabertura(tmp_path, monkeypatch):
    caminho = tmp_path / "vendas.xlsx"
    respostas = iter(["Caneta", "3", "2,50", "s", "Caderno", "1", "18", "n"])
    monkeypatch.setattr("builtins.input", lambda *_: next(respostas))
    aut.main(caminho)

    ws = load_workbook(caminho)[aut.NOME_ABA]
    assert [ws["A2"].value, ws["B2"].value, ws["C2"].value] == ["Caneta", 3, 2.5]
    assert ws["A3"].value == "Caderno"

    # segunda execução continua de onde parou
    respostas = iter(["Borracha", "5", "1", "n"])
    monkeypatch.setattr("builtins.input", lambda *_: next(respostas))
    aut.main(caminho)
    ws = load_workbook(caminho)[aut.NOME_ABA]
    assert ws["A4"].value == "Borracha"
