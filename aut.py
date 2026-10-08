"""AUT-EX: automação de planilha de vendas em Excel.

Cria (ou abre) a planilha vendas.xlsx, registra produtos em sequência,
deixa o Excel calcular os totais com fórmulas e mantém o visual organizado.

Uso:
    python aut.py
"""
import math
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ARQUIVO = Path("vendas.xlsx")
NOME_ABA = "Vendas"
CABECALHO = ["Produto", "Quantidade", "Preço", "Total"]
FORMATO_MOEDA = '"R$" #,##0.00'
COR_CABECALHO = "1F4E78"


# ---------------------------------------------------------------- planilha

def estilizar_cabecalho(ws):
    for celula in ws[1][: len(CABECALHO)]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
        celula.alignment = Alignment(horizontal="center", vertical="center")


def criar_planilha():
    wb = Workbook()
    ws = wb.active
    ws.title = NOME_ABA
    ws.append(CABECALHO)
    return wb, ws


def abrir_planilha(caminho):
    """Abre a planilha existente ou cria uma nova. Retorna (workbook, aba)."""
    try:
        wb = load_workbook(caminho)
    except FileNotFoundError:
        print("Nova planilha criada!")
        return criar_planilha()
    except BadZipFile:
        raise SystemExit(f"'{caminho}' não parece ser um arquivo Excel válido. Renomeie ou apague e tente de novo.")

    ws = wb[NOME_ABA] if NOME_ABA in wb.sheetnames else wb.active
    print("Planilha carregada!")
    return wb, ws


def ultima_linha_dados(ws):
    """Última linha com produto na coluna A (1 se só existir o cabeçalho)."""
    linha = ws.max_row
    while linha > 1 and ws.cell(linha, 1).value is None:
        linha -= 1
    return linha


def adicionar_venda(ws, produto, quantidade, preco):
    """Escreve uma venda na próxima linha livre. O total é uma fórmula do Excel."""
    linha = ultima_linha_dados(ws) + 1
    ws.cell(linha, 1, produto)
    ws.cell(linha, 2, quantidade)
    ws.cell(linha, 3, preco)
    ws.cell(linha, 4, f"=B{linha}*C{linha}")
    return linha


def formatar_planilha(ws):
    """Aplica estilo, formatos, filtro, painel congelado, resumo e larguras."""
    ultima = ultima_linha_dados(ws)

    estilizar_cabecalho(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{ultima}"

    for linha in range(2, ultima + 1):
        ws.cell(linha, 3).number_format = FORMATO_MOEDA
        ws.cell(linha, 4).number_format = FORMATO_MOEDA

    # resumo ao lado da tabela (SUM ignora o texto do cabeçalho)
    ws["F1"], ws["G1"] = "Total geral", "=SUM(D:D)"
    ws["F2"], ws["G2"] = "Itens vendidos", "=SUM(B:B)"
    ws["F1"].font = ws["F2"].font = Font(bold=True)
    ws["G1"].number_format = FORMATO_MOEDA

    produtos = [len(str(ws.cell(l, 1).value)) for l in range(2, ultima + 1)]
    larguras = {"A": min(max([len("Produto")] + produtos) + 4, 50), "B": 14, "C": 16, "D": 18, "E": 3, "F": 18, "G": 18}
    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura


def salvar(wb, caminho):
    """Salva o arquivo; se estiver aberto no Excel, espera o usuário fechar."""
    while True:
        try:
            wb.save(caminho)
            return
        except PermissionError:
            input(f"Não consegui salvar. Feche '{caminho}' no Excel e aperte Enter para tentar de novo... ")


# ---------------------------------------------------------------- entrada

def converter_numero(texto, tipo, minimo):
    """Converte '10', '10.5', '10,5' ou '1.500,50' em número validado.

    Levanta ValueError se não for número ou estiver abaixo do mínimo.
    """
    texto = texto.strip()
    if "," in texto:  # formato brasileiro: ponto separa milhar, vírgula separa decimal
        texto = texto.replace(".", "").replace(",", ".")
    valor = tipo(texto)
    if not math.isfinite(valor):
        raise ValueError("número inválido")
    if valor < minimo:
        raise ValueError(f"o valor mínimo é {minimo}")
    return valor


def ler_texto(mensagem):
    while True:
        valor = input(mensagem).strip()
        if valor:
            return valor
        print("  Esse campo é obrigatório.")


def ler_numero(mensagem, tipo, minimo):
    while True:
        try:
            return converter_numero(input(mensagem), tipo, minimo)
        except ValueError as erro:
            print(f"  Valor inválido ({erro}). Tente de novo.")


def quer_continuar():
    resposta = input("Adicionar outro produto? [S/n] ").strip().lower()
    return resposta not in ("n", "nao", "não")


# ---------------------------------------------------------------- programa

def main(caminho=ARQUIVO):
    wb, ws = abrir_planilha(caminho)
    adicionadas = 0

    try:
        while True:
            produto = ler_texto("Produto: ")
            quantidade = ler_numero("Quantidade: ", int, 1)
            preco = round(ler_numero("Preço: ", float, 0), 2)

            linha = adicionar_venda(ws, produto, quantidade, preco)
            adicionadas += 1
            print(f"  OK: {produto} | {quantidade} x R$ {preco:.2f} = R$ {quantidade * preco:.2f} (linha {linha})")

            if not quer_continuar():
                break
    except (KeyboardInterrupt, EOFError):
        print("\nEncerrando...")

    if adicionadas:
        formatar_planilha(ws)
        salvar(wb, caminho)
        print(f"{adicionadas} produto(s) salvo(s) em '{caminho}'.")
    else:
        print("Nada novo para salvar.")


if __name__ == "__main__":
    main()
